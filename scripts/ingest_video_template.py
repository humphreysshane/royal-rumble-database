# -*- coding: utf-8 -*-
"""
Ingests a completed Video_Analysis_Entry_Template.xlsx (built by
build_video_template.py, in the project root) into entrances.csv /
near_eliminations.csv / moves.csv.

    python3 ingest_video_template.py <path-to-filled-in.xlsx> <data-dir>

APPENDS rows -- never overwrites what's already in these tables. Always run
this against an ISOLATED COPY of data/ first (per this project's standing
protocol), then re-run validate_integrity.py and eyeball the new rows,
before running it again against the live data/ directory.

What this script does NOT do: it does not run build_derived.py or
build_dashboard_data.py, and none of these three tables feed into either
script yet (they're hand-populated source tables, not wired into any
derived stat or dashboard page as of 2026-09-23 -- see ROADMAP.md /
IDEAS.md "Phase C"). Surfacing this data on the dashboard is a separate,
later step once there's enough of it to be worth a UI.
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from schema import ENTRANCES_FIELDS, NEAR_ELIMINATIONS_FIELDS, MOVES_FIELDS

from openpyxl import load_workbook

SHEETS = [
    ("Entrances", "entrances.csv", ENTRANCES_FIELDS),
    ("Near Eliminations", "near_eliminations.csv", NEAR_ELIMINATIONS_FIELDS),
    ("Moves", "moves.csv", MOVES_FIELDS),
]

# The exact example rows build_video_template.py writes -- skipped even if
# Shane forgets to delete them, as a safety net (the Read Me also tells him
# to delete them before sending the file back).
EXAMPLE_SENTINELS = {
    "Entrances": ("RR2020M", "drew-mcintyre"),
    "Near Eliminations": ("RR2018W", "sasha-banks"),
    "Moves": ("RR2020M", "brock-lesnar"),
}


def cellstr(v):
    """Normalize an openpyxl cell value back to this database's string
    convention: TRUE/FALSE for Excel booleans, '' for blank, str() for
    anything else (Excel may hand back an int/float for a numeric-looking
    column like entrance_duration_seconds)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 ingest_video_template.py <filled-in.xlsx> <data-dir>")
        sys.exit(2)
    xlsx_path, data_dir = sys.argv[1], sys.argv[2]

    wb = load_workbook(xlsx_path, data_only=True)

    # Cross-check ids against the live database rather than trusting the
    # sheet blindly -- catches a typo'd event_id/wrestler_id before it's
    # silently written as a new, wrong row.
    with open(os.path.join(data_dir, "events.csv"), newline="", encoding="utf-8") as f:
        known_events = {r["event_id"] for r in csv.DictReader(f)}
    with open(os.path.join(data_dir, "wrestlers.csv"), newline="", encoding="utf-8") as f:
        known_wrestlers = {r["wrestler_id"] for r in csv.DictReader(f)}
    id_fields_by_sheet = {
        "Entrances": ["wrestler_id"],
        "Near Eliminations": ["wrestler_id", "attempted_by_id", "pulled_back_in_by_id"],
        "Moves": ["wrestler_performing_id", "wrestler_receiving_id"],
    }

    problems = []
    totals = {}

    for sheet_name, csv_name, fields in SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"WARNING: sheet {sheet_name!r} missing from workbook, skipping.")
            continue
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        if header != fields:
            problems.append(f"{sheet_name}: header row doesn't match schema.py's "
                             f"{csv_name} field list -- don't rename/reorder columns. "
                             f"Expected {fields}, got {header}")
            continue

        new_rows = []
        sentinel = EXAMPLE_SENTINELS[sheet_name]
        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            values = [cellstr(c.value) for c in row]
            if not any(values):
                continue  # fully blank row, skip silently
            rowdict = dict(zip(fields, values))
            if sheet_name == "Entrances" and (rowdict.get("event_id"), rowdict.get("wrestler_id")) == sentinel:
                continue
            if sheet_name == "Near Eliminations" and (rowdict.get("event_id"), rowdict.get("wrestler_id")) == sentinel:
                continue
            if sheet_name == "Moves" and (rowdict.get("event_id"), rowdict.get("wrestler_performing_id")) == sentinel:
                continue

            eid = rowdict.get("event_id", "")
            if eid and eid not in known_events:
                problems.append(f"{sheet_name} row {r_idx}: event_id {eid!r} not found in events.csv")
            for f in id_fields_by_sheet[sheet_name]:
                wid = rowdict.get(f, "")
                if wid and wid not in known_wrestlers:
                    problems.append(f"{sheet_name} row {r_idx}: {f}={wid!r} not found in wrestlers.csv")

            new_rows.append(rowdict)

        totals[csv_name] = new_rows

        # append to the CSV (create with header if it doesn't exist / is empty)
        csv_path = os.path.join(data_dir, csv_name)
        file_exists = os.path.exists(csv_path) and os.path.getsize(csv_path) > 0
        with open(csv_path, "a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            if not file_exists:
                writer.writeheader()
            writer.writerows(new_rows)

    if problems:
        print(f"PROBLEMS FOUND ({len(problems)}) -- rows were still appended for any sheet with a "
              f"correct header, but review these before treating this as clean:")
        for p in problems:
            print(" -", p)

    for csv_name, rows in totals.items():
        print(f"{csv_name}: appended {len(rows)} row(s)")


if __name__ == "__main__":
    main()
